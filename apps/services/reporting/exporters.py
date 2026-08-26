"""
Exportadores del Reporting Hub -- FASE 21 (mision Reporting Hub).

Todos consumen un ReportResult ya calculado -- NINGUNO vuelve a consultar la
base de datos (Query -> ReportResult -> {JSON, CSV, XLSX, ...}).

PDF queda explicitamente fuera de esta pasada: no hay libreria de PDF
instalada en requirements.txt (reportlab/weasyprint ausentes) y agregar una
es una decision de dependencia que corresponde confirmar, no asumir --
ver docs/reporting/REPORTING_BASELINE.md §5.
"""
from __future__ import annotations

import csv
import io

from apps.services.reporting.contracts import ReportResult

SUPPORTED_FORMATS = ("json", "csv", "xlsx")


def export_json(result: ReportResult) -> dict:
    return result.to_dict()


def export_csv(result: ReportResult) -> bytes:
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=list(result.columns), extrasaction="ignore")
    writer.writeheader()
    for row in result.rows:
        writer.writerow(row)
    if result.totals:
        totals_row = {col: result.totals.get(col, "") for col in result.columns}
        if result.dimensions:
            totals_row[result.dimensions[0]] = "TOTAL"
        writer.writerow(totals_row)
    return buffer.getvalue().encode("utf-8-sig")


def export_xlsx(result: ReportResult) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = result.dataset_id[:31] or "Reporte"

    ws.append(list(result.columns))
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in result.rows:
        ws.append([row.get(col, "") for col in result.columns])

    if result.totals:
        totals_row = [result.totals.get(col, "") for col in result.columns]
        if result.dimensions:
            totals_row[result.columns.index(result.dimensions[0])] = "TOTAL"
        ws.append(totals_row)
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)

    for col_cells in ws.columns:
        length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=10)
        ws.column_dimensions[col_cells[0].column_letter].width = min(length + 2, 40)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


_EXPORTERS = {
    "csv": export_csv,
    "xlsx": export_xlsx,
}


def export(result: ReportResult, fmt: str) -> bytes | dict:
    fmt = fmt.lower()
    if fmt == "json":
        return export_json(result)
    exporter = _EXPORTERS.get(fmt)
    if exporter is None:
        raise ValueError(f"Formato de exportacion no soportado: {fmt}")
    return exporter(result)
